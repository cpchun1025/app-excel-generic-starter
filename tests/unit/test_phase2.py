"""Portable Phase 2 tests using generated, non-sensitive workbooks only."""

import hashlib
import logging
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from excel_automation.cli import main
from excel_automation.config.validation import load_and_validate
from excel_automation.errors import ExecutionError, OutputCollisionError
from excel_automation.runner.direct_services import DirectXlsxServices
from excel_automation.runner.orchestrator import run_job
from excel_automation.samples import generate_sample
from excel_automation.workbook.copies import stage_copy
from excel_automation.workbook.direct_xlsx import DirectXlsx, bounds
from excel_automation.workbook.validation import validate_output


@pytest.fixture
def plain_job(tmp_path):
    source = tmp_path / "input" / "plain.xlsx"
    generate_sample(source)
    config = tmp_path / "job.yaml"
    config.write_text(
        yaml.safe_dump({
            "schema_version": 1,
            "jobs": {
                "plain": {
                    "inputs": {"source": "input/plain.xlsx"},
                    "workbook": {
                        "base": "source",
                        "output": "output/result.xlsx",
                        "adapter": "direct_xlsx",
                    },
                    "steps": [{
                        "id": "check-output",
                        "type": "validate",
                        "checks": [
                            {"type": "file_exists"},
                            {"type": "sheet_exists", "sheet": "Sample"},
                            {"type": "range_nonempty", "sheet": "Sample", "range": "A1:C3"},
                            {
                                "type": "range_equals", "sheet": "Sample", "range": "B2",
                                "expected": 12,
                            },
                        ],
                    }],
                },
            },
        }),
        encoding="utf-8",
    )
    return load_and_validate(config).jobs["plain"], config


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_staging_copies_bytes_without_touching_source(plain_job, tmp_path):
    job, _ = plain_job
    source = job.inputs[job.base]
    before = digest(source)
    staged = tmp_path / "staged.xlsx"
    stage_copy(source, staged)
    assert digest(staged) == before
    assert digest(source) == before
    with pytest.raises(ExecutionError):
        stage_copy(source, staged)
    assert digest(source) == before
    with pytest.raises(ExecutionError):
        stage_copy(source, source)


def test_plain_values_roundtrip_on_staged_copy(plain_job, tmp_path):
    job, _ = plain_job
    source = job.inputs[job.base]
    before = digest(source)
    staged = tmp_path / "staged.xlsx"
    stage_copy(source, staged)
    values = [["Text & <XML> \"quotes\"", 3.25, False], ["=literal text", 0, None]]
    with DirectXlsx(staged) as adapter:
        assert adapter.read_values("Sample", "A2:C2") == [["Fake pencils", 12, True]]
        adapter.write_values("Sample", "A2:C3", values)
        adapter.save()
    with DirectXlsx(staged) as adapter:
        assert adapter.read_values("Sample", "A2:C3") == values
    assert digest(source) == before


@pytest.mark.parametrize("dry_run", [False, True])
def test_real_run_source_unchanged_and_publication_policy(plain_job, dry_run):
    job, _ = plain_job
    source = job.inputs[job.base]
    before = digest(source)
    contexts = []

    class ObserveStaging(DirectXlsxServices):
        def execute(self, step, context):
            contexts.append(context)
            assert context.staged_output.is_file()
            assert context.staged_output != job.output
            assert not job.output.exists()
            assert context.fake_execution is False
            return super().execute(step, context)

    result = run_job(job, ObserveStaging(), dry_run=dry_run)
    assert result["status"] == "real_execution_succeeded"
    assert result["fake_execution"] is False
    assert result["output_published"] is not dry_run
    assert digest(source) == before
    assert all(not context.workspace.exists() for context in contexts)
    if dry_run:
        assert not job.output.exists()
        assert not job.output.parent.exists()
    else:
        assert digest(job.output) == before
        assert list(job.output.parent.iterdir()) == [job.output]


@pytest.mark.parametrize("dry_run", [False, True])
def test_real_collision_preserves_existing_output(plain_job, dry_run):
    job, _ = plain_job
    job.output.parent.mkdir()
    job.output.write_bytes(b"previous output")
    with pytest.raises(OutputCollisionError):
        run_job(job, DirectXlsxServices(), dry_run=dry_run)
    assert job.output.read_bytes() == b"previous output"


@pytest.mark.parametrize("kind", ["python", "copy", "macro", "email", "render"])
def test_unimplemented_steps_fail_before_staging(plain_job, kind):
    job, _ = plain_job
    unsupported = replace(job.steps[0], id="unimplemented", type=kind)
    job = replace(job, steps=(*job.steps, unsupported))

    class NoStage(DirectXlsxServices):
        def stage(self, job, context):
            pytest.fail("Unsupported step must be rejected before staging")

    with pytest.raises(ExecutionError, match="Unimplemented real step"):
        run_job(job, NoStage())
    assert not job.output.parent.exists()


def test_original_examples_are_schema_valid_but_not_real_executable(monkeypatch):
    from excel_automation.plugins.example_allowlist import APPROVED_EXAMPLE_PLUGIN_IDS

    root = Path(__file__).resolve().parents[2]
    config = load_and_validate(
        root / "config" / "jobs.example.yaml",
        plugin_ids=APPROVED_EXAMPLE_PLUGIN_IDS,
    )
    for job in config.jobs.values():
        assert DirectXlsxServices.capability_error(job)
        with pytest.raises(ExecutionError):
            run_job(job, DirectXlsxServices())


@pytest.mark.parametrize("failure", ["stage", "validation"])
def test_failure_cleans_staging_without_completed_output(plain_job, failure):
    job, _ = plain_job
    contexts = []

    class FailingServices(DirectXlsxServices):
        def stage(self, job, context):
            contexts.append(context)
            super().stage(job, context)
            if failure == "stage":
                raise ExecutionError("Injected stage failure")

        def execute(self, step, context):
            context.staged_output.unlink()
            return super().execute(step, context)

    with pytest.raises(ExecutionError):
        run_job(job, FailingServices())
    assert not job.output.exists()
    assert all(not context.workspace.exists() for context in contexts)


def test_validation_failure_prevents_publication(plain_job):
    job, _ = plain_job
    bad_step = replace(
        job.steps[0],
        definition={"checks": [{"type": "sheet_exists", "sheet": "Missing"}]},
    )
    with pytest.raises(ExecutionError, match="worksheet is missing"):
        run_job(replace(job, steps=(bad_step,)), DirectXlsxServices())
    assert not job.output.exists()


def test_publication_race_never_overwrites_existing_output(plain_job, monkeypatch):
    from excel_automation.workbook.copies import publish_output

    job, _ = plain_job

    def concurrent_publication(staged, output):
        output.write_bytes(b"concurrent completed output")
        publish_output(staged, output)

    monkeypatch.setattr(
        "excel_automation.runner.orchestrator.publish_output", concurrent_publication
    )
    with pytest.raises(OutputCollisionError):
        run_job(job, DirectXlsxServices())
    assert job.output.read_bytes() == b"concurrent completed output"
    assert list(job.output.parent.iterdir()) == [job.output]


def test_publication_filesystem_failure_has_no_final_output(plain_job, monkeypatch):
    job, _ = plain_job

    def no_links(*args):
        raise OSError("unsupported filesystem")

    monkeypatch.setattr("excel_automation.workbook.copies.os.link", no_links)
    with pytest.raises(ExecutionError, match="hard-link support"):
        run_job(job, DirectXlsxServices())
    assert not job.output.exists()
    assert list(job.output.parent.iterdir()) == []


@pytest.mark.parametrize(
    "reference",
    ["A0", "XFE1", "A1048577", "B2:A1", "A1:XFD1048576", "Sheet1!A1", "A:A"],
)
def test_range_boundaries_are_checked(reference):
    with pytest.raises(ExecutionError):
        bounds(reference)


def test_exact_cell_and_package_limits(plain_job, monkeypatch):
    from excel_automation.workbook import direct_xlsx

    assert bounds("A1:A100000") == (1, 1, 1, 100000)
    with pytest.raises(ExecutionError, match="100000-cell"):
        bounds("A1:A100001")
    job, _ = plain_job
    source = job.inputs[job.base]
    with ZipFile(source) as package:
        actual_size = sum(item.file_size for item in package.infolist())
    monkeypatch.setattr(direct_xlsx, "MAX_PACKAGE_BYTES", actual_size)
    with DirectXlsx(source):
        pass
    monkeypatch.setattr(direct_xlsx, "MAX_PACKAGE_BYTES", actual_size - 1)
    with pytest.raises(ExecutionError, match="package limit"):
        DirectXlsx(source)


def test_checks_on_empty_zero_false_and_expected_values(plain_job):
    job, _ = plain_job
    source = job.inputs[job.base]
    for reference in ("C3", "B2"):
        validate_output(source, [
            {"type": "range_nonempty", "sheet": "Sample", "range": reference}
        ])
    with pytest.raises(ExecutionError, match="range is empty"):
        validate_output(source, [
            {"type": "range_nonempty", "sheet": "Sample", "range": "D1:D2"}
        ])
    with pytest.raises(ExecutionError, match="values omitted"):
        validate_output(source, [{
            "type": "range_equals", "sheet": "Sample", "range": "A2",
            "expected": "sensitive-value-not-to-log",
        }])
    with DirectXlsx(source) as adapter:
        with pytest.raises(ExecutionError, match="does not exist"):
            adapter.read_values("Missing", "A1")
        with pytest.raises(ExecutionError, match="dimensions"):
            adapter.write_values("Sample", "A1:B1", [[1]])
        with pytest.raises(ExecutionError, match="Only plain"):
            adapter.write_values("Sample", "A1", [[datetime(2025, 1, 1)]])
        with pytest.raises(ExecutionError, match="Only plain"):
            adapter.write_values("Sample", "A1", [[float("nan")]])
        with pytest.raises(ExecutionError, match="Only plain"):
            adapter.write_values("Sample", "A1", [[10**1000]])


@pytest.mark.parametrize("payload", [b"", b"not-a-workbook", b"PK malformed archive"])
def test_malformed_input_fails_before_staging(plain_job, payload):
    job, _ = plain_job
    job.inputs[job.base].write_bytes(payload)
    with pytest.raises(ExecutionError, match="Cannot inspect workbook"):
        run_job(job, DirectXlsxServices())
    assert not job.output.parent.exists()


def test_missing_input_and_empty_output_report_errors(plain_job, tmp_path):
    job, _ = plain_job
    job.inputs[job.base].unlink()
    with pytest.raises(ExecutionError):
        run_job(job, DirectXlsxServices())
    empty = tmp_path / "empty.xlsx"
    empty.touch()
    with pytest.raises(ExecutionError, match="missing or empty"):
        validate_output(empty, [{"type": "file_exists"}])


@pytest.mark.parametrize(
    "feature", [
        "formula", "style", "date", "merge", "hyperlink", "table", "chart",
        "hidden", "protection", "conditional_formatting", "data_validation",
    ]
)
def test_known_unsupported_workbook_features_rejected_before_staging(plain_job, feature):
    job, _ = plain_job
    source = job.inputs[job.base]
    workbook = load_workbook(source)
    sheet = workbook.active
    if feature == "formula":
        sheet["D1"] = "=SUM(B2:B3)"
    elif feature == "style":
        sheet["A1"].font = Font(bold=True)
    elif feature == "date":
        sheet["D1"] = datetime(2025, 1, 1)
    elif feature == "merge":
        sheet.merge_cells("D1:E1")
    elif feature == "hyperlink":
        sheet["A2"].hyperlink = "https://example.invalid"
    elif feature == "table":
        sheet.add_table(Table(displayName="FakeTable", ref="A1:C3"))
    elif feature == "chart":
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        sheet.add_chart(chart, "E1")
    elif feature == "hidden":
        workbook.create_sheet("Hidden").sheet_state = "hidden"
    elif feature == "protection":
        sheet.protection.sheet = True
    elif feature == "conditional_formatting":
        sheet.conditional_formatting.add("B2:B3", CellIsRule(operator="greaterThan", formula=["1"]))
    elif feature == "data_validation":
        validation = DataValidation(type="whole", operator="between", formula1=1, formula2=20)
        sheet.add_data_validation(validation)
        validation.add("B2:B3")
    workbook.save(source)
    workbook.close()
    before = digest(source)
    with pytest.raises(ExecutionError, match="outside the supported"):
        run_job(job, DirectXlsxServices())
    assert digest(source) == before
    assert not job.output.parent.exists()


@pytest.mark.parametrize("extra_part", [
    "xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml",
    "xl/pivotTables/pivotTable1.xml", "xl/drawings/drawing1.xml",
    "xl/sharedStrings.xml", "customXml/item1.xml",
])
def test_unknown_package_parts_rejected(plain_job, extra_part):
    job, _ = plain_job
    source = job.inputs[job.base]
    with ZipFile(source, "a") as package:
        package.writestr(extra_part, b"<unsupported/>")
    with pytest.raises(ExecutionError, match="outside the supported"):
        run_job(job, DirectXlsxServices())
    assert not job.output.parent.exists()


def test_multiple_plain_sheets_empty_sheet_and_duplicate_values(tmp_path):
    path = tmp_path / "multi.xlsx"
    workbook = Workbook()
    workbook.active.append(["duplicate", "duplicate", 0, False])
    workbook.create_sheet("Empty")
    workbook.save(path)
    workbook.close()
    with DirectXlsx(path) as adapter:
        assert adapter.read_values("Sheet", "A1:D1") == [["duplicate", "duplicate", 0, False]]
        assert adapter.read_values("Empty", "A1") == [[None]]
        adapter.save()
    with DirectXlsx(path) as adapter:
        assert adapter.has_sheet("Empty")


def test_empty_workbook_is_supported_but_nonempty_check_fails(tmp_path):
    path = tmp_path / "empty.xlsx"
    workbook = Workbook()
    workbook.save(path)
    workbook.close()
    validate_output(path, [{"type": "file_exists"}, {"type": "sheet_exists", "sheet": "Sheet"}])
    with pytest.raises(ExecutionError, match="range is empty"):
        validate_output(path, [{"type": "range_nonempty", "sheet": "Sheet", "range": "A1"}])


def test_malformed_xml_and_duplicate_parts_are_rejected(plain_job, tmp_path):
    job, _ = plain_job
    source = job.inputs[job.base]
    with ZipFile(source) as package:
        contents = {name: package.read(name) for name in package.namelist()}
    malformed = tmp_path / "malformed.xlsx"
    with ZipFile(malformed, "w") as package:
        for name, data in contents.items():
            package.writestr(name, b"<not-closed" if name.endswith("sheet1.xml") else data)
    with pytest.raises(ExecutionError, match="Cannot inspect"):
        DirectXlsx(malformed)
    duplicate = tmp_path / "duplicate.xlsx"
    with ZipFile(duplicate, "w") as package:
        for name, data in contents.items():
            package.writestr(name, data)
        with pytest.warns(UserWarning, match="Duplicate name"):
            package.writestr("xl/styles.xml", contents["xl/styles.xml"])
    with pytest.raises(ExecutionError, match="outside the supported"):
        DirectXlsx(duplicate)


def test_xlsm_extension_is_rejected_without_modification(plain_job, tmp_path):
    job, _ = plain_job
    path = tmp_path / "renamed.xlsm"
    path.write_bytes(job.inputs[job.base].read_bytes())
    before = digest(path)
    with pytest.raises(ExecutionError, match="only .xlsx"):
        DirectXlsx(path)
    assert digest(path) == before


@pytest.mark.parametrize("dry_run", [False, True])
def test_real_execution_has_no_email_transport(plain_job, monkeypatch, dry_run):
    import smtplib

    def forbidden_smtp(*args, **kwargs):
        pytest.fail("Phase 2 must never contact SMTP")

    monkeypatch.setattr(smtplib, "SMTP", forbidden_smtp)
    monkeypatch.setattr(smtplib, "SMTP_SSL", forbidden_smtp)
    job, _ = plain_job
    run_job(job, DirectXlsxServices(), dry_run=dry_run)


def test_invalid_check_shapes_rejected_before_staging(plain_job):
    job, _ = plain_job
    for expected in ([1, 2], [[1, 2]], []):
        step = replace(job.steps[0], definition={"checks": [{
            "type": "range_equals", "sheet": "Sample", "range": "A1", "expected": expected,
        }]})
        with pytest.raises(ExecutionError, match="rectangular matrix"):
            run_job(replace(job, steps=(step,)), DirectXlsxServices())
        assert not job.output.parent.exists()


def test_numeric_equality_excludes_boolean_coercion(plain_job):
    job, _ = plain_job
    source = job.inputs[job.base]
    validate_output(source, [{
        "type": "range_equals", "sheet": "Sample", "range": "B2", "expected": 12.0,
    }])
    with pytest.raises(ExecutionError, match="values omitted"):
        validate_output(source, [{
            "type": "range_equals", "sheet": "Sample", "range": "C2", "expected": 1,
        }])


def test_phase2_cli_real_run_and_dry_run(plain_job, capsys):
    job, config = plain_job
    assert main(["--config", str(config), "run", job.id, "--dry-run"]) == 0
    output = capsys.readouterr().out
    assert "REAL DIRECT_XLSX RUN SUCCEEDED" in output
    assert "final output not published" in output
    assert "FAKE" not in output
    assert not job.output.exists()
    assert main(["--config", str(config), "run", job.id]) == 0
    assert job.output.is_file()
    assert "output published" in capsys.readouterr().out


def test_cli_selects_eligible_job_without_running_ineligible_sibling(plain_job, capsys):
    job, config = plain_job
    document = yaml.safe_load(config.read_text(encoding="utf-8"))
    sibling = {
        "inputs": {"source": "not-present.xlsx"},
        "workbook": {
            "base": "source", "output": "not-created.xlsx", "adapter": "direct_xlsx",
        },
        "steps": [{
            "id": "plugin", "type": "python", "plugin": "sample.clean_sales",
        }],
    }
    document["jobs"]["ineligible"] = sibling
    config.write_text(yaml.safe_dump(document), encoding="utf-8")
    assert main(["--config", str(config), "list-jobs"]) == 0
    listing = capsys.readouterr().out
    assert "real-direct_xlsx=eligible" in listing
    assert "production-executable=no" in listing
    assert main(["--config", str(config), "run", job.id, "--dry-run"]) == 0
    assert "REAL DIRECT_XLSX RUN SUCCEEDED" in capsys.readouterr().out
    assert main(["--config", str(config), "run", "ineligible"]) == 3
    assert "Unimplemented real step" in capsys.readouterr().err
    assert not (config.parent / "not-created.xlsx").exists()


def test_sample_cli_without_config_and_no_overwrite(tmp_path, capsys):
    output = tmp_path / "fake.xlsx"
    assert main(["generate-samples", "--output", str(output)]) == 0
    before = digest(output)
    assert main(["generate-samples", "--output", str(output)]) == 3
    assert digest(output) == before
    assert "no output overwritten" in capsys.readouterr().err


def test_real_failure_logs_no_cell_values(plain_job, caplog):
    job, _ = plain_job
    step = replace(job.steps[0], definition={"checks": [{
        "type": "range_equals", "sheet": "Sample", "range": "A2",
        "expected": "CONFIDENTIAL-EXPECTED-CELL",
    }]})
    logger = logging.getLogger("phase2_test")
    with caplog.at_level(logging.INFO, logger=logger.name):
        with pytest.raises(ExecutionError):
            run_job(replace(job, steps=(step,)), DirectXlsxServices(), logger=logger)
    assert "status=failed" in caplog.text
    assert "CONFIDENTIAL" not in caplog.text
    assert "Fake pencils" not in caplog.text
