"""Plain-value .xlsx operations; deliberately not a general fidelity adapter."""

import math
import re
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import fromstring
from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils.cell import get_column_letter, range_boundaries

from excel_automation.errors import ExecutionError

CellValue = str | int | float | bool | None
MAX_CELLS = 100_000
MAX_PACKAGE_BYTES = 20_000_000
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A1 = re.compile(r"^\$?[A-Z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Z]{1,3}\$?[1-9][0-9]*)?$")


def bounds(reference: str) -> tuple[int, int, int, int]:
    """Validate bounded, non-reversed A1 ranges before allocating cells."""
    if not isinstance(reference, str) or not A1.fullmatch(reference):
        raise ExecutionError("Expected a bounded A1 cell or range")
    left, top, right, bottom = range_boundaries(reference)
    if not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1048576):
        raise ExecutionError("Range is reversed or outside Excel worksheet bounds")
    if (right - left + 1) * (bottom - top + 1) > MAX_CELLS:
        raise ExecutionError("Range exceeds the direct_xlsx 100000-cell limit")
    return left, top, right, bottom


def _shape(element):
    return (
        element.tag,
        sorted(element.attrib.items()),
        (element.text or "").strip(),
        tuple(_shape(child) for child in element),
    )


@lru_cache(maxsize=1)
def _baseline() -> dict[str, bytes]:
    workbook = Workbook()
    try:
        stream = BytesIO()
        workbook.save(stream)
        with ZipFile(stream) as package:
            return {name: package.read(name) for name in package.namelist()}
    finally:
        workbook.close()


def _reject() -> None:
    raise ExecutionError(
        "Workbook is outside the supported plain-value direct_xlsx subset; "
        "no workbook was rewritten"
    )


def inspect_package(path: Path) -> None:
    """Reject unknown parts, non-default features, and non-plain worksheet XML."""
    if path.suffix.lower() != ".xlsx":
        raise ExecutionError("direct_xlsx accepts only .xlsx workbooks")
    try:
        with ZipFile(path) as package:
            names = package.namelist()
            baseline = _baseline()
            if len(names) != len(set(names)):
                _reject()
            if sum(item.file_size for item in package.infolist()) > MAX_PACKAGE_BYTES:
                raise ExecutionError("Workbook exceeds the 20 MB uncompressed package limit")
            allowed = set(baseline) - {"xl/worksheets/sheet1.xml"}
            sheets = [name for name in names if re.fullmatch(r"xl/worksheets/sheet[1-9][0-9]*\.xml", name)]
            if not sheets or set(names) != allowed | set(sheets):
                _reject()
            roots = {name: fromstring(package.read(name)) for name in names}
            for name in ("xl/styles.xml", "xl/theme/theme1.xml"):
                if _shape(roots[name]) != _shape(fromstring(baseline[name])):
                    _reject()
            workbook = roots["xl/workbook.xml"]
            default_workbook = fromstring(baseline["xl/workbook.xml"])
            if workbook.tag != default_workbook.tag or workbook.attrib or [
                _shape(child) for child in workbook if child.tag != f"{{{MAIN}}}sheets"
            ] != [
                _shape(child) for child in default_workbook if child.tag != f"{{{MAIN}}}sheets"
            ]:
                _reject()
            sheet_nodes = workbook.findall(f"{{{MAIN}}}sheets/{{{MAIN}}}sheet")
            containers = workbook.findall(f"{{{MAIN}}}sheets")
            if (
                len(sheet_nodes) != len(sheets)
                or len(containers) != 1
                or containers[0].attrib
                or list(containers[0]) != sheet_nodes
                or len({node.get("name", "").casefold() for node in sheet_nodes}) != len(sheet_nodes)
            ):
                _reject()
            for sheet in sheet_nodes:
                if set(sheet.attrib) != {"name", "sheetId", "state", f"{{{REL}}}id"}:
                    _reject()
                if sheet.attrib["state"] != "visible" or list(sheet):
                    _reject()
            # Relationships/content types are limited to the same generated
            # package profile, with only the worksheet count allowed to vary.
            for name in ("_rels/.rels", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                root = roots[name]
                expected = fromstring(baseline[name])
                if root.tag != expected.tag or root.attrib:
                    _reject()
                for child in root:
                    if list(child):
                        _reject()
                    if _shape(child) in [_shape(node) for node in expected]:
                        continue
                    attrs = child.attrib
                    if name == "xl/_rels/workbook.xml.rels":
                        normalized = dict(attrs)
                        normalized.pop("Id", None)
                        if any(
                            child.tag == node.tag
                            and normalized == {key: value for key, value in node.attrib.items() if key != "Id"}
                            for node in expected
                            if node.get("Type") != f"{REL}/worksheet"
                        ):
                            continue
                        if (
                            child.tag != expected[0].tag
                            or set(attrs) != {"Type", "Target", "Id"}
                            or attrs["Type"] != f"{REL}/worksheet"
                            or attrs["Target"].lstrip("/") not in sheets
                        ):
                            _reject()
                    elif name == "[Content_Types].xml":
                        if (
                            child.tag != expected[-1].tag
                            or set(attrs) != {"PartName", "ContentType"}
                            or attrs["PartName"].lstrip("/") not in sheets
                            or attrs["ContentType"] != "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
                        ):
                            _reject()
                    else:
                        _reject()
            default_sheet = fromstring(baseline["xl/worksheets/sheet1.xml"])
            for name in sheets:
                root = roots[name]
                variable = {f"{{{MAIN}}}sheetData", f"{{{MAIN}}}dimension"}
                if root.tag != default_sheet.tag or root.attrib or [
                    _shape(child) for child in root if child.tag not in variable
                ] != [_shape(child) for child in default_sheet if child.tag not in variable]:
                    _reject()
                dimensions = root.findall(f"{{{MAIN}}}dimension")
                data = root.findall(f"{{{MAIN}}}sheetData")
                if len(dimensions) != 1 or len(data) != 1:
                    _reject()
                if set(dimensions[0].attrib) != {"ref"} or list(dimensions[0]) or data[0].attrib:
                    _reject()
                bounds(dimensions[0].get("ref", ""))
                seen_cells: set[str] = set()
                seen_rows: set[int] = set()
                for row in data[0]:
                    if row.tag != f"{{{MAIN}}}row" or set(row.attrib) != {"r"}:
                        _reject()
                    row_number = int(row.attrib["r"])
                    if not 1 <= row_number <= 1048576 or row_number in seen_rows:
                        _reject()
                    seen_rows.add(row_number)
                    for cell in row:
                        if cell.tag != f"{{{MAIN}}}c" or set(cell.attrib) - {"r", "t"}:
                            _reject()
                        reference = cell.get("r", "")
                        _, cell_row, _, _ = bounds(reference)
                        if reference in seen_cells or ":" in reference or cell_row != row_number:
                            _reject()
                        seen_cells.add(reference)
                        if len(seen_cells) > MAX_CELLS:
                            _reject()
                        cell_type = cell.get("t", "n")
                        if cell_type not in {"n", "b", "inlineStr"}:
                            _reject()
                        if cell_type == "inlineStr":
                            if len(cell) == 0:  # openpyxl serializes empty strings this way
                                continue
                            if len(cell) != 1 or cell[0].tag != f"{{{MAIN}}}is" or cell[0].attrib:
                                _reject()
                            if len(cell[0]) != 1 or cell[0][0].tag != f"{{{MAIN}}}t":
                                _reject()
                            if list(cell[0][0]) or set(cell[0][0].attrib) - {"{http://www.w3.org/XML/1998/namespace}space"}:
                                _reject()
                        elif len(cell) > 1 or any(
                            node.tag != f"{{{MAIN}}}v" or node.attrib or list(node) for node in cell
                        ):
                            _reject()
    except (OSError, BadZipFile, ParseError, DefusedXmlException, ValueError, KeyError) as exc:
        raise ExecutionError(
            f"Cannot inspect workbook package ({type(exc).__name__}); contents omitted"
        ) from exc


class DirectXlsx:
    """Context-managed read/write access to a vetted, staged plain workbook."""

    def __init__(self, path: Path) -> None:
        inspect_package(path)
        self.path = path
        try:
            self.workbook = load_workbook(path, data_only=False, keep_links=False)
        except (OSError, ValueError, KeyError, TypeError, IndexError, BadZipFile) as exc:
            raise ExecutionError(
                f"Cannot open plain workbook ({type(exc).__name__}); contents omitted"
            ) from exc
        try:
            for sheet in self.workbook:
                bounds(sheet.calculate_dimension())
                for row in sheet:
                    for cell in row:
                        if cell.data_type in {"f", "e", "d"} or not _plain_value(cell.value):
                            _reject()
        except ExecutionError:
            self.workbook.close()
            raise

    def __enter__(self) -> "DirectXlsx":
        return self

    def __exit__(self, *args) -> None:
        self.workbook.close()

    def has_sheet(self, name: str) -> bool:
        return name in self.workbook.sheetnames

    def _sheet(self, name: str):
        if not self.has_sheet(name):
            raise ExecutionError("Required worksheet does not exist; cell contents omitted")
        return self.workbook[name]

    def read_values(self, sheet: str, reference: str) -> list[list[CellValue]]:
        left, top, right, bottom = bounds(reference)
        worksheet = self._sheet(sheet)
        bounds(f"A1:{get_column_letter(max(worksheet.max_column, right))}{max(worksheet.max_row, bottom)}")
        values = [
            [worksheet.cell(row, col).value for col in range(left, right + 1)]
            for row in range(top, bottom + 1)
        ]
        return values

    def write_values(self, sheet: str, reference: str, values: list[list[CellValue]]) -> None:
        left, top, right, bottom = bounds(reference)
        if len(values) != bottom - top + 1 or any(len(row) != right - left + 1 for row in values):
            raise ExecutionError("Write matrix dimensions must match the target range")
        if any(not _plain_value(value) for row in values for value in row):
            raise ExecutionError("Only plain text, finite numbers, booleans, and blank values are supported")
        worksheet = self._sheet(sheet)
        new_right = max(worksheet.max_column, right)
        new_bottom = max(worksheet.max_row, bottom)
        bounds(f"A1:{get_column_letter(new_right)}{new_bottom}")
        for row_index, row in enumerate(values, top):
            for col_index, value in enumerate(row, left):
                cell = worksheet.cell(row_index, col_index, value)
                if value is None:
                    cell.value = None
                elif isinstance(value, str):
                    cell.data_type = "s"

    def save(self) -> None:
        try:
            self.workbook.save(self.path)
        except (OSError, ValueError, TypeError) as exc:
            raise ExecutionError(
                f"Cannot save staged workbook ({type(exc).__name__}); contents omitted"
            ) from exc


def _plain_value(value: object) -> bool:
    if value is None or type(value) is bool:
        return True
    if type(value) is int:
        return abs(value) <= 10**308
    if type(value) is float:
        return math.isfinite(value)
    if isinstance(value, str):
        return len(value) <= 32767 and not ILLEGAL_CHARACTERS_RE.search(value)
    return False
